import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from contextlib import ExitStack
import os
import shutil
from openpyxl import load_workbook
from datetime import date
from breeders_cup.updates import upsert,save_state,merge_research

ROOT=Path(__file__).resolve().parents[1]


class UpdateTests(unittest.TestCase):
    def setUp(self):
        self.data=json.loads((ROOT/'examples/demo.json').read_text())
        self.template=ROOT/'templates/2024 Breeders Cup Final.xlsx'

    def test_create_then_refresh_preserves_picks_results_and_summary(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'tests') as tmp:
            target=Path(tmp)/'generated/Breeders Cup 2026.xlsx'
            initial,_=upsert(self.template,self.data,target)
            w=load_workbook(target);s=w['2026']
            picked=next(r for r in s if s.cell(r[0].row,19).value=='demo-alpha' and s.cell(r[0].row,20).value=='demo-classic')
            s.cell(picked[0].row,5).value='DE'
            s.cell(picked[0].row,3).value=2
            s.cell(picked[0].row,2).value=4
            summary=next(r[0].row for r in s if r[0].value=='RESULTS SUMMARY')
            s.cell(summary+1,2).value=7
            w['Overage']['E2']='H'
            w.save(target);w.close()
            newer=copy.deepcopy(self.data)
            newer['horses'].append({'id':'new','name':'A New Horse','source_url':'test'})
            newer['entries'].append({'horse_id':'new','race_id':'demo-classic','post':1,'odds':'5-1'})
            next(e for e in newer['entries'] if e['horse_id']=='demo-alpha' and e['race_id']=='demo-classic').update(scratched=True,post='SCR')
            updated,_=upsert(self.template,newer,target,previous=initial)
            w=load_workbook(target);s=w['2026']
            row=next(r for r in s if s.cell(r[0].row,19).value=='demo-alpha' and s.cell(r[0].row,20).value=='demo-classic')
            self.assertEqual([row[c-1].value for c in (2,3,4,5)],[4,2,'SCR','DE'])
            self.assertTrue(any(r[4].value=='H' and r[5].value=='Demo Alpha' for r in w['Overage']))
            header=next(r[0].row for r in s if r[0].value=='RESULTS SUMMARY')
            self.assertEqual(s.cell(header+1,2).value,7)
            w.close()
            save_state(updated,Path(tmp)/'data')
            self.assertEqual([p.name for p in target.parent.iterdir()],['Breeders Cup 2026.xlsx'])
            self.assertTrue((Path(tmp)/'data/state/2026.json').exists())
            self.assertTrue((Path(tmp)/'data/reports/2026.json').exists())

    def test_disappearing_picked_entry_is_retained_not_inferred_scratched(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'tests') as tmp:
            target=Path(tmp)/'Breeders Cup 2026.xlsx'
            prior,_=upsert(self.template,self.data,target)
            w=load_workbook(target)
            row=next(r for r in w['2026'] if w['2026'].cell(r[0].row,19).value=='demo-beta')
            w['2026'].cell(row[0].row,5).value='H';w.save(target);w.close()
            new=copy.deepcopy(self.data)
            new['entries']=[e for e in new['entries'] if e['horse_id']!='demo-beta']
            new['horses']=[h for h in new['horses'] if h['id']!='demo-beta']
            _,_=upsert(self.template,new,target,previous=prior)
            w=load_workbook(target)
            row=next(r for r in w['2026'] if w['2026'].cell(r[0].row,19).value=='demo-beta')
            self.assertEqual(row[4].value,'H');self.assertNotEqual(row[3].value,'SCR')
            w.close()

    def test_research_reuse_rejects_changed_parent_identity(self):
        old=copy.deepcopy(self.data)
        old['horses'][0]['grandfathers']='Known A; Known B'
        self.assertEqual(merge_research(self.data,old)['horses'][0]['grandfathers'],'Known A; Known B')
        new=copy.deepcopy(self.data);new['horses'][0]['father']='Different sire'
        self.assertNotIn('grandfathers',merge_research(new,old)['horses'][0])

    def test_default_cli_recreates_generated_and_keeps_one_workbook(self):
        from breeders_cup.cli import main
        with tempfile.TemporaryDirectory(dir=ROOT/'tests') as tmp,ExitStack() as stack:
            folder=Path(tmp)
            (folder/'templates').mkdir()
            shutil.copyfile(self.template,folder/'templates/2024 Breeders Cup Final.xlsx')
            original_cwd=Path.cwd()
            stack.callback(os.chdir,original_cwd)
            os.chdir(folder)
            stack.enter_context(patch('sys.argv',['breeders-cup']))
            stack.enter_context(patch('breeders_cup.cli.current_year',return_value=2026))
            stack.enter_context(patch('breeders_cup.cli.collect',side_effect=lambda year:copy.deepcopy(self.data)))
            for name in ('repair_incomplete_identities','enrich','enrich_remaining','enrich_query','annotate_famous','refresh_jockeys'):
                stack.enter_context(patch('breeders_cup.cli.'+name,side_effect=lambda data,*args:data))
            self.assertEqual(main(),0)
            target=folder/'generated/Breeders Cup 2026.xlsx'
            w=load_workbook(target);w['2026']['E3']='DE';w.save(target);w.close()
            self.assertEqual(main(),0)
            self.assertEqual([p.name for p in (folder/'generated').iterdir()],['Breeders Cup 2026.xlsx'])
            w=load_workbook(target);self.assertEqual(w['2026']['E3'].value,'DE');w.close()
            self.assertTrue((folder/'data/state/2026.json').exists())

    def test_saturday_preserves_friday_rows_and_summary_but_refreshes_saturday(self):
        data=copy.deepcopy(self.data)
        data['races'][0]['date']='2026-10-30'
        for i,entry in enumerate(data['entries']):
            entry.update(post=i+1,odds='5-1')
        with tempfile.TemporaryDirectory(dir=ROOT/'tests') as tmp:
            target=Path(tmp)/'Breeders Cup 2026.xlsx'
            prior,_=upsert(self.template,data,target,as_of=date(2026,10,30))
            w=load_workbook(target);s=w['2026']
            friday=next(r for r in s if s.cell(r[0].row,20).value=='demo-turf')
            row=friday[0].row
            s.cell(row,5).value='DE';s.cell(row,3).value=1
            s.cell(row,7).value='Saved Friday Rider'
            header=next(r[0].row for r in s if r[0].value=='RESULTS SUMMARY')
            s.cell(header+1,2).value=1;s.cell(header+1,3).value=2
            before=[s.cell(row,c).value for c in range(1,21)]
            old_summary=[s.cell(header+1,c).value for c in range(1,7)]
            w.save(target);w.close()
            fresh=copy.deepcopy(data)
            fresh['horses'][0]['jockey']='Changed Saturday Rider'
            fresh['entries'][0].update(post=9,odds='50-1',scratched=True)
            fresh['entries'][1].update(post=4,odds='12-1')
            # A newly reported Friday entrant cannot alter an archived field.
            fresh['entries'].append({'horse_id':'demo-beta','race_id':'demo-turf','post':2,'odds':'2-1'})
            updated,_=upsert(self.template,fresh,target,previous=prior,as_of=date(2026,10,31))
            w=load_workbook(target);s=w['2026']
            friday=[r for r in s if s.cell(r[0].row,20).value=='demo-turf']
            self.assertEqual(len(friday),1)
            self.assertEqual([s.cell(friday[0][0].row,c).value for c in range(1,21)],before)
            header=next(r[0].row for r in s if r[0].value=='RESULTS SUMMARY')
            self.assertEqual([s.cell(header+1,c).value for c in range(1,7)],old_summary)
            saturday=next(r for r in s if s.cell(r[0].row,20).value=='demo-classic' and s.cell(r[0].row,19).value=='demo-alpha')
            self.assertEqual(saturday[3].value,'12-1')
            self.assertEqual(saturday[6].value,'Changed Saturday Rider')
            self.assertEqual(len(updated['collection_report']['yearly_update']['past_races_preserved']),1)
            w.close()
            # Even disappearing unpicked Friday rows are retained.
            fresh['entries']=[e for e in fresh['entries'] if e['race_id']!='demo-turf']
            fresh['races']=[r for r in fresh['races'] if r['id']!='demo-turf']
            upsert(self.template,fresh,target,previous=updated,as_of=date(2026,10,31))
            w=load_workbook(target)
            self.assertTrue(any(w['2026'].cell(r[0].row,20).value=='demo-turf' for r in w['2026']))
            w.close()

    def test_same_day_stays_live_and_friday_only_overage_is_frozen_next_day(self):
        data=copy.deepcopy(self.data)
        data['races'][0]['date']='2026-10-30'
        data['entries']=[e for e in data['entries'] if not(e['horse_id']=='demo-alpha' and e['race_id']=='demo-classic')]
        data['entries'][0]['odds']='5-1'
        with tempfile.TemporaryDirectory(dir=ROOT/'tests') as tmp:
            target=Path(tmp)/'Breeders Cup 2026.xlsx'
            prior,_=upsert(self.template,data,target,as_of=date(2026,10,30))
            fresh=copy.deepcopy(data);fresh['entries'][0]['odds']='8-1'
            prior,_=upsert(self.template,fresh,target,previous=prior,as_of=date(2026,10,30))
            w=load_workbook(target);self.assertEqual(w['2026']['D3'].value,'8-1');w.close()
            fresh['entries'][0]['odds']='50-1';fresh['horses'][0]['trainer']='New trainer'
            upsert(self.template,fresh,target,previous=prior,as_of=date(2026,10,31))
            w=load_workbook(target)
            self.assertEqual(w['2026']['D3'].value,'8-1')
            self.assertEqual(w['Overage']['D2'].value,'8-1')
            self.assertNotEqual(w['Overage']['H2'].value,'New trainer')
            w.close()
