import copy
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from breeders_cup.core import prepare
from breeders_cup.workbook import generate

ROOT = Path(__file__).resolve().parents[1]


class YearlyTests(unittest.TestCase):
    def test_early_research_without_schedule_and_future_year_isolation(self):
        data = json.loads((ROOT / 'examples/early-2026-demo.json').read_text())
        self.assertEqual(prepare(data)[2], [])
        with self.assertRaisesRegex(ValueError, 'assignment required'):
            prepare(data, final=True)
        (ROOT / 'generated').mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=ROOT / 'generated') as tmp:
            first = Path(tmp) / 'Breeders Cup 2026.xlsx'
            generate(ROOT / 'templates/2024 Breeders Cup Final.xlsx', data, first)
            unchanged = first.read_bytes()
            w = load_workbook(first)
            self.assertEqual(w.sheetnames, ['2026', 'Overage'])
            self.assertEqual(w['2026'].max_row, 1)
            self.assertEqual([w['Overage'].cell(r, 6).value for r in range(2, 5)],
                             ['Demo Alpha', 'Demo Beta', 'Demo Charlie'])
            self.assertTrue(all(w['Overage'].cell(r, 1).value is None for r in range(2, 5)))
            self.assertEqual(w['Overage']['I2'].value, 'Bay')
            w.close()
            data['year'] = 2027
            second = Path(tmp) / 'Breeders Cup 2027.xlsx'
            generate(ROOT / 'templates/2024 Breeders Cup Final.xlsx', data, second)
            w = load_workbook(second)
            self.assertEqual(w.sheetnames, ['2027', 'Overage'])
            w.close()
            self.assertEqual(first.read_bytes(), unchanged)

    def test_assigned_and_unassigned_horses_remain_in_overage(self):
        data = json.loads((ROOT / 'examples/demo.json').read_text())
        data['horses'].append({'id': 'pending', 'name': 'Pending Horse', 'source_url': 'Synthetic'})
        data['assignments'] = {'demo-alpha': 'demo-classic'}
        (ROOT / 'generated').mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=ROOT / 'generated') as tmp:
            target = Path(tmp) / 'assigned.xlsx'
            self.assertEqual(generate(ROOT / 'templates/2024 Breeders Cup Final.xlsx', data, target), 2)
            w = load_workbook(target)
            self.assertEqual(w['Overage'].max_row, 4)
            self.assertEqual(w['Overage']['A2'].value, 'Classic')
            self.assertEqual(w['Overage']['F4'].value, 'Pending Horse')
            self.assertIsNone(w['Overage']['A4'].value)
            w.close()
