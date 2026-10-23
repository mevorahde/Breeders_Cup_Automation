import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from breeders_cup.core import prepare, last_race_before
from breeders_cup.workbook import generate

ROOT = Path(__file__).resolve().parents[1]


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.data = json.loads((ROOT / 'examples/demo.json').read_text())

    def test_final_cross_entry_requires_assignment(self):
        with self.assertRaisesRegex(ValueError, 'assignment required'):
            prepare(self.data, final=True)
        self.data['assignments'] = {'demo-alpha': 'demo-classic'}
        _, _, entries = prepare(self.data, final=True)
        self.assertEqual(len(entries), 2)
        self.assertFalse(any(e['race_id'] == 'demo-turf' for e in entries))

    def test_order_is_year_specific(self):
        self.data['races'].reverse()
        self.assertEqual([r['number'] for r in prepare(self.data)[0]], [8, 9])
        self.data['races'][0]['date'] = '2026-10-30'
        with self.assertRaisesRegex(ValueError, 'Saturday'):
            prepare(self.data)

    def test_invalid_and_duplicate_entries(self):
        self.data['entries'].append(copy.deepcopy(self.data['entries'][0]))
        with self.assertRaises(ValueError):
            prepare(self.data)

    def test_previous_race_excludes_future_and_uncompleted(self):
        history = [{'date': '2026-09-01', 'name': 'Prior'},
                   {'date': '2026-10-31', 'name': 'Cup'},
                   {'date': '2026-10-01', 'name': 'Cancelled', 'status': 'cancelled'}]
        self.assertEqual(last_race_before(history, '2026-10-31')['name'], 'Prior')

    def test_workbook_layout_overage_and_no_overwrite(self):
        template = ROOT / 'templates/2024 Breeders Cup Final.xlsx'
        self.assertTrue(template.exists(), 'Reference template must be installed')
        self.data['horses'][1]['name'] = '=DEMO formula-looking text'
        self.data['horses'][0]['pedigree_ancestors'] = ['Secretariat']
        self.data['entries'][1].update(post=10, pick='DE')
        self.data['entries'][2].update(post=2, pick='H')
        (ROOT / 'generated').mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=ROOT / 'generated') as tmp:
            target = Path(tmp) / 'demo.xlsx'
            with patch('breeders_cup.workbook.FAMOUS_HORSES', ['Secretariat']):
                self.assertEqual(generate(template, self.data, target), 3)
            w = load_workbook(target)
            self.assertEqual(w.sheetnames, ['2026', 'Overage'])
            self.assertEqual(w['Overage'].max_row, 3)
            self.assertEqual(w['Overage']['F2'].data_type, 's')
            self.assertEqual(w['2026']['A2'].value, 'Saturday 10-31-2026')
            self.assertEqual(w['2026']['O3'].value.year, 2022)
            self.assertEqual([w['2026'][f'B{r}'].value for r in (5, 6)], [2, 10])
            self.assertIsNone(w['2026']['E5'].value)
            self.assertIsNone(w['2026']['E6'].value)
            self.assertEqual(w['2026']['N3'].value, 'Secretariat')
            self.assertGreater(w['2026'].column_dimensions['F'].width, 11.7265625)
            self.assertEqual(w['2026'].column_dimensions['L'].width, w['Overage'].column_dimensions['L'].width)
            self.assertEqual(w['2026'].row_dimensions[3].height, 12)
            reference = load_workbook(template)
            self.assertEqual(copy.copy(w['2026']['F3'].font), copy.copy(reference['2024']['F5'].font))
            reference.close()
            w.close()
            with self.assertRaises(FileExistsError):
                generate(template, self.data, target)

    def test_test_generation_caps_each_race_and_keeps_overage(self):
        seed=copy.deepcopy(self.data['horses'][0])
        seed['jockey']='A Confirmed Rider With A Long Name'
        self.data['horses']=[]
        self.data['entries']=[]
        for i in range(18):
            horse={**seed,'id':f'horse-{i}','name':f'Horse {i:02}'}
            self.data['horses'].append(horse)
            for race in self.data['races']:
                self.data['entries'].append({'horse_id':horse['id'],'race_id':race['id'],'post':18-i})
        (ROOT/'generated').mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=ROOT/'generated') as tmp:
            target=Path(tmp)/'test.xlsx'
            self.assertEqual(generate(ROOT/'templates/2024 Breeders Cup Final.xlsx',self.data,target,test_generation=True),26)
            w=load_workbook(target)
            self.assertEqual(w['Overage'].max_row,19)
            self.assertEqual(w['Overage']['G2'].value,seed['jockey'])
            self.assertGreater(w['Overage'].column_dimensions['G'].width,25)
            self.assertTrue(all(w['Overage'].column_dimensions[c].bestFit for c in 'ABCDEFGHIJKLMNO'))
            s=w['2026']
            start=next(r[0].row for r in s if r[0].value=='RESULTS SUMMARY')
            for race in self.data['races']:
                rows=[r for r in s.iter_rows(max_row=start-2) if r[0].value==race['name']]
                self.assertEqual(len(rows),13)
                self.assertEqual([r[1].value for r in rows],list(range(1,14)))
            self.assertEqual(s.cell(start+1,1).value,prepare(self.data)[0][0]['name'])
            self.assertTrue(all(s.cell(r,c).value is None for r in range(start+1,start+3) for c in range(2,6)))
            self.assertEqual(s.cell(start+3,1).value,'SubTotals')
            self.assertIn(f'C{start+1}:C{start+2}',s.cell(start+3,3).value)
            self.assertEqual(s.cell(start+4,1).value,'Low Total')
            self.assertEqual(s.cell(start+5,1).value,'Final Totals')
            for col in ('C','D','E'):
                self.assertLess(s.column_dimensions[col].width,15,'Formula text must not widen result columns')
            w.close()
            self.assertEqual(generate(ROOT/'templates/2024 Breeders Cup Final.xlsx',self.data,Path(tmp)/'normal.xlsx'),36)
            with self.assertRaisesRegex(ValueError,'cannot be combined'):
                generate(ROOT/'templates/2024 Breeders Cup Final.xlsx',self.data,Path(tmp)/'invalid.xlsx',final=True,test_generation=True)


if __name__ == '__main__':
    unittest.main()
