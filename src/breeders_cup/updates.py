"""One yearly workbook, with stable row matching and atomic pick-preserving saves."""
from copy import deepcopy,copy
from datetime import datetime,timezone,timedelta
from pathlib import Path
import hashlib
import json
import os
import tempfile
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from .pedigree import normalized
from .core import prepare
from .workbook import generate,fit_columns


def current_year():
    try:
        from zoneinfo import ZoneInfo
        zone=ZoneInfo('America/Los_Angeles')
    except (ImportError,KeyError):
        # The LA calendar-year boundary is always in standard time (UTC-8).
        zone=timezone(timedelta(hours=-8))
    return datetime.now(zone).year


def merge_research(data,previous):
    result=deepcopy(data)
    if not previous or previous.get('year')!=result['year']:
        return result
    old={h['id']:h for h in previous['horses']}
    fields=('grandfathers','great_grandfathers','pedigree_ancestors','pedigree_branches',
            'famous_ancestor_evidence','other_ancestors','date_of_birth','color','father','dam',
            'maternal_grandsire','identity_source','identity_method')
    for horse in result['horses']:
        prior=old.get(horse['id'])
        if not prior or normalized(prior['name'])!=normalized(horse['name']):
            continue
        if any(horse.get(f) not in (None,'','--') and prior.get(f) not in (None,'','--') and
               normalized(horse[f])!=normalized(prior[f]) for f in ('father','dam','date_of_birth')):
            continue
        for field in fields:
            if not horse.get(field) or horse.get(field)=='--':
                if prior.get(field) and prior[field]!='--':
                    horse[field]=deepcopy(prior[field])
                    if field in prior.get('field_sources',{}):
                        horse.setdefault('field_sources',{})[field]=prior['field_sources'][field]
        if not horse.get('jockey') and prior.get('jockey_status')=='confirmed':
            for field in ('jockey','jockey_status','jockey_retrieved_at'):
                if prior.get(field):
                    horse[field]=prior[field]
            if prior.get('field_sources',{}).get('jockey'):
                horse.setdefault('field_sources',{})['jockey']=prior['field_sources']['jockey']
    return result


def race_rows(sheet):
    stop=next((r[0].row for r in sheet if r[0].value=='RESULTS SUMMARY'),sheet.max_row+1)
    return [r for r in sheet.iter_rows(min_row=2,max_row=stop-1) if r[5].value]


def capture(book,data,previous):
    """Read manual input before rebuilding. IDs are preferred; names migrate old files."""
    races={normalized(r['name']):r['id'] for r in data['races']}
    horses={normalized(h['name']):h['id'] for h in data['horses']}
    if previous:
        races.update({normalized(r['name']):r['id'] for r in previous['races']})
        horses.update({normalized(h['name']):h['id'] for h in previous['horses']})
    horse_names={}
    race_names={}
    for records in (data,previous or data):
        for h in records['horses']:
            horse_names.setdefault(h['id'],set()).add(normalized(h['name']))
        for r in records['races']:
            race_names.setdefault(r['id'],set()).add(normalized(r['name']))
    saved={}
    for sn in (str(data['year']),'Overage'):
        if sn not in book:
            raise ValueError(f'Existing workbook is missing {sn}; it was not changed')
        sheet=book[sn]
        rows=race_rows(sheet) if sn!='Overage' else list(sheet.iter_rows(min_row=2))
        for row in rows:
            if not row[5].value:
                continue
            hid=row[18].value if len(row)>18 else None
            if hid and normalized(row[5].value) not in horse_names.get(hid,set()):
                hid=None  # Recover safely after a manual sort excluding hidden IDs.
            hid=hid or horses.get(normalized(row[5].value))
            rid=None if sn=='Overage' else (row[19].value if len(row)>19 else None) or races.get(normalized(row[0].value))
            if rid and normalized(row[0].value) not in race_names.get(rid,set()):
                rid=races.get(normalized(row[0].value))
            if not hid or (sn!='Overage' and not rid):
                if row[4].value is not None or row[2].value is not None:
                    raise ValueError(f'Cannot safely identify a row containing picks/results: {sn}!F{row[5].row}')
                continue
            key=(sn,hid,rid)
            if key in saved:
                raise ValueError('Duplicate saved horse/race identity; workbook was not changed')
            saved[key]={c:copy(row[c-1]) for c in (2,3,4,5)}
    return saved


def preserve_summary(old,new):
    old_header=next((r[0].row for r in old if r[0].value=='RESULTS SUMMARY'),None)
    new_header=next((r[0].row for r in new if r[0].value=='RESULTS SUMMARY'),None)
    if old_header is None or new_header is None:
        return
    old_rows={normalized(old.cell(r,1).value):r for r in range(old_header+1,old.max_row+1) if old.cell(r,1).value}
    for r in range(new_header+1,new.max_row+1):
        prior=old_rows.get(normalized(new.cell(r,1).value))
        if prior:
            for c in range(2,6):
                cell=old.cell(prior,c)
                # Generated totals are recreated with new row references.
                if cell.value is not None and cell.data_type!='f':
                    new.cell(r,c).value=cell.value


def upsert(template,data,output,previous=None,final=False,test_generation=False):
    data=deepcopy(data)
    output=Path(output)
    if data.get('demo'):
        raise ValueError('DEMO data cannot update the production yearly workbook')
    if output.resolve()==Path(template).resolve():
        raise ValueError('Output must differ from template')
    old=None
    fingerprint=None
    saved={}
    if output.exists():
        fingerprint=hashlib.sha256(output.read_bytes()).hexdigest()
        old=load_workbook(output)
        if str(data['year']) not in old:
            old.close()
            raise ValueError('Existing workbook belongs to a different year')
        try:
            saved=capture(old,data,previous)
        except Exception:
            old.close()
            raise
    old_horses={h['id']:h for h in (previous or data)['horses']}
    old_races={r['id']:r for r in (previous or data)['races']}
    horse_ids={h['id'] for h in data['horses']}
    race_ids={r['id'] for r in data['races']}
    pairs={(e['horse_id'],e['race_id']) for e in data['entries']}
    for (sn,hid,rid),cells in saved.items():
        if sn=='Overage' and hid not in horse_ids and any(cells[c].value is not None for c in (3,5)):
            if hid not in old_horses:
                raise ValueError('Saved Overage pick has no previous research identity; workbook was not changed')
            data['horses'].append(deepcopy(old_horses[hid]));horse_ids.add(hid)
        if sn=='Overage' or (hid,rid) in pairs or not any(cells[c].value is not None for c in (3,5)):
            continue
        # A disappearing source row must not silently discard someone's pick/result.
        if hid not in old_horses or rid not in old_races:
            raise ValueError('Saved pick has no previous research identity; workbook was not changed')
        if hid not in horse_ids:
            data['horses'].append(deepcopy(old_horses[hid]));horse_ids.add(hid)
        if rid not in race_ids:
            data['races'].append(deepcopy(old_races[rid]));race_ids.add(rid)
        data['entries'].append({'horse_id':hid,'race_id':rid,'retained_manual_row':True})
        pairs.add((hid,rid))
    for entry in data['entries']:
        if entry.get('scratched') or str(entry.get('post','')).upper()=='SCR':
            entry['odds']='SCR'
            prior=saved.get((str(data['year']),entry['horse_id'],entry['race_id']))
            if str(entry.get('post','')).upper()=='SCR' and prior and str(prior[2].value).isdigit():
                entry['post']=prior[2].value
    # Explicit withdrawal notices also cover horses removed from the live cards.
    by_horse={h['id']:h for h in data['horses']}
    by_race={r['id']:r for r in data['races']}
    for line in data.get('withdrawals',[]):
        if ' has withdrawn from the ' not in line:
            continue
        name,race=line.split(' has withdrawn from the ',1)
        race=race.split(' (AE')[0]
        for entry in data['entries']:
            if normalized(by_horse[entry['horse_id']]['name'])==normalized(name) and normalized(by_race[entry['race_id']]['name'])==normalized(race):
                entry.update(scratched=True,odds='SCR')
    prepare(data,final)
    output.parent.mkdir(parents=True,exist_ok=True)
    try:
        with tempfile.TemporaryDirectory(dir=output.parent,prefix='.updating-') as directory:
            temporary=Path(directory)/output.name
            count=generate(template,data,temporary,final,test_generation)
            book=load_workbook(temporary)
            restored=set()
            for sn in (str(data['year']),'Overage'):
                sheet=book[sn]
                rows=race_rows(sheet) if sn!='Overage' else list(sheet.iter_rows(min_row=2))
                for row in rows:
                    hid=row[18].value if len(row)>18 else None
                    rid=None if sn=='Overage' else row[19].value
                    prior=saved.get((sn,hid,rid))
                    if not prior:
                        continue
                    restored.add((sn,hid,rid))
                    for col in (3,5):
                        source=prior[col]
                        if source.value is not None:
                            cell=sheet.cell(row[0].row,col)
                            cell.value=Translator(source.value,origin=source.coordinate).translate_formula(cell.coordinate) if source.data_type=='f' else source.value
                            if col==5:
                                cell.font=copy(source.font);cell.fill=copy(source.fill)
                    for col in (2,4):
                        cell=sheet.cell(row[0].row,col)
                        # Missing source data must not erase the last known post/odds.
                        if cell.value is None:
                            cell.value=prior[col].value
                if old and sn!='Overage':
                    preserve_summary(old[sn],sheet)
            if any(key not in restored and any(cells[c].value is not None for c in (3,5)) for key,cells in saved.items()):
                book.close()
                raise ValueError('Update would hide a saved pick/result; workbook was not changed')
            fit_columns(book[str(data['year'])],book['Overage'])
            book.save(temporary);book.close()
            if old:
                old.close();old=None
            if fingerprint is not None:
                if not output.exists() or hashlib.sha256(output.read_bytes()).hexdigest()!=fingerprint:
                    raise ValueError('Workbook changed during update; rerun to preserve the latest edits')
                os.replace(temporary,output)
            else:
                # Exclusive creation prevents overwriting a concurrent first run.
                with output.open('xb') as stream:
                    stream.write(temporary.read_bytes())
    finally:
        if old:
            old.close()
    data.setdefault('collection_report',{})['yearly_update']={'file':str(output),'updated_existing':fingerprint is not None,
        'displayed_race_entries':count,'picks_preserved':sum(cells[5].value is not None for cells in saved.values())}
    return data,count


def save_state(data,root=Path('data')):
    for folder,value in [('state',data),('reports',data.get('collection_report',{}))]:
        target=Path(root)/folder/f'{data["year"]}.json'
        target.parent.mkdir(parents=True,exist_ok=True)
        with tempfile.NamedTemporaryFile(mode='w',encoding='utf8',dir=target.parent,delete=False,suffix='.tmp') as stream:
            json.dump(value,stream,indent=2,ensure_ascii=False)
            temporary=Path(stream.name)
        os.replace(temporary,target)
