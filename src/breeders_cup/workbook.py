"""Create a new workbook from the user's 2024 styles; never overwrite an output."""
from copy import copy
from datetime import date
from pathlib import Path
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.formula.translate import Translator
from .core import FIELDS, prepare, famous_ancestors, post_order
from .famous_horses import FAMOUS_HORSES
from .column_fit import fitted_width

LEGACY_HEADERS = ['BC Race', 'Post', 'Result', 'Odds', 'Pick', 'Name', 'Jockey',
           'Trainer', 'Color', 'Last Previous Race', 'Father', 'Grandfather(s)',
           'Great Grandfather(s)', 'Other Ancestors', 'Date of Birth']
HEADERS = LEGACY_HEADERS[:5] + ['Date of Birth'] + LEGACY_HEADERS[5:14]


def arrange_columns(book, legacy=False):
    """Move birth dates for display; retain the template layout during refresh.

    Summary cells stay in their established positions. Hidden source/identity
    columns are never moved. Both layouts can be read for existing workbooks.
    """
    desired = LEGACY_HEADERS if legacy else HEADERS
    for sheet in book:
        current = [sheet.cell(1, col).value for col in range(1, 16)]
        if current == desired:
            continue
        if current not in (LEGACY_HEADERS, HEADERS) and current[1:] not in (LEGACY_HEADERS[1:], HEADERS[1:]):
            raise ValueError('Unrecognized workbook columns; workbook was not changed')
        # A demo may label the first header BC Race (DEMO).
        current[0] = 'BC Race'
        mapping = {col: current.index(label) + 1 for col, label in enumerate(desired, 1)}
        stop = next((r[0].row for r in sheet if r[0].value == 'RESULTS SUMMARY'), sheet.max_row + 1)
        for row in range(1, stop):
            cells = {col: copy(sheet.cell(row, col)) for col in range(6, 16)}
            for col in range(6, 16):
                source = cells[mapping[col]]
                target = sheet.cell(row, col)
                # Set a private style before assigning dates: openpyxl may change
                # a General format on assignment, otherwise mutating shared styles.
                target._style = copy(source._style)
                target.value = (Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
                                if source.data_type == 'f' else source.value)
                target.data_type = source.data_type
                target.comment = copy(source.comment)
                target.hyperlink = copy(source.hyperlink)
        sheet.freeze_panes = 'F2' if legacy else 'G2'


def title_case_details(value):
    """Title-case display text, retaining country codes and name punctuation."""
    if not isinstance(value, str):
        return value
    countries = {'USA', 'GB', 'IRE', 'FR', 'GER', 'JPN', 'AUS', 'NZ', 'CAN',
                 'ARG', 'BRZ', 'CHI', 'SAF', 'ITY', 'UAE', 'KSA', 'URU', 'PER'}
    def word(match):
        token = match.group()
        if token in {'SCR', 'AE'}:
            return token
        # Preserve established internal capitals such as McKinzie and Le Havre.
        if not (token.isupper() or token.islower()):
            return token
        titled = token.title()
        return re.sub(r"(['’])S$", r"\1s", titled)
    result = re.sub(r"[^\W\d_]+(?:['’][^\W\d_]+)*", word, value)
    return re.sub(r'\(([A-Za-z]+)\)',
                  lambda m: '(' + (m[1].upper() if m[1].upper() in countries else m[1]) + ')',
                  result)


def display_color(value):
    """Use the user's compact coat-color labels in the workbook."""
    if not isinstance(value, str):
        return value
    return {'dark bay or brown': 'DK Bay', 'gray or roan': 'Gray',
            'dk bay': 'DK Bay'}.get(value.strip().casefold(), value)


def display_last_race(value):
    """Keep the race and finish, omitting dates and periods."""
    if not isinstance(value, str):
        return value
    return re.sub(r'\s*\(\s*\d{4}-\d{2}-\d{2}\s*\)', '', value).replace('.', '').strip()


def text(cell, value):
    # Scraped text must remain text, including formula-looking horse names.
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'


def styled_row(target, row, values, source, source_row):
    target.row_dimensions[row].height = source.row_dimensions[source_row].height
    for col, value in enumerate(values, 1):
        if source_row != 1 and 6 <= col <= 14:
            value = title_case_details(value)
        if col == 9:
            value = display_color(value)
        if col == 10:
            value = display_last_race(value)
        cell = target.cell(row, col)
        original = source.cell(source_row, min(col, 14))
        cell.font = copy(original.font)
        cell.fill = copy(original.fill)
        cell.border = copy(original.border)
        cell.number_format = original.number_format
        cell.protection = copy(original.protection)
        cell.alignment = Alignment(horizontal='center' if col <= 5 or col in (9, 15) else 'left',
                                   vertical='center', wrap_text=False)
        text(cell, value)
        if isinstance(value, date):
            cell.number_format = 'mm/dd/yyyy'
    target.row_dimensions[row].height = source.row_dimensions[source_row].height or 12


def results_summary(sheet, row, races, source):
    header = next((r[0].row for r in source.iter_rows(min_col=1,max_col=1)
                   if r[0].value=='RESULTS SUMMARY'),None)
    if header is None:
        raise ValueError('Template is missing RESULTS SUMMARY')
    def copy_style(cell,original):
        for attr in ('font','fill','border','alignment','protection'):
            setattr(cell,attr,copy(getattr(original,attr)))
        cell.number_format=original.number_format
    def copy_summary_row(destination, template_row, values):
        for col in range(1,16):
            cell=sheet.cell(destination,col)
            original=source.cell(template_row,col)
            copy_style(cell,original)
            if col<=len(values):
                value=values[col-1]
                if isinstance(value,str) and value.startswith('='):
                    cell.value=value
                else:
                    text(cell,value)
        sheet.row_dimensions[destination].height=12
    # Keep the template's black separator above the summary.
    copy_summary_row(row,header-1,[])
    for cell in sheet[row]:
        cell.fill=copy(source.cell(header-1,1).fill)
    row+=1
    copy_summary_row(row,header,['RESULTS SUMMARY','De','$','H','$','Win = $2'])
    first=row+1
    for i,race in enumerate(races):
        copy_summary_row(first+i,header+1,[race['name'],None,None,None,None,
            'Place = $1' if i==0 else 'Low Total = $10' if i==1 else None])
        if i>1:
            copy_style(sheet.cell(first+i,6),source.cell(header+4,6))
    last=first+len(races)-1
    subtotal=last+1
    low=last+2
    final=last+3
    copy_summary_row(subtotal,header+15,['SubTotals'])
    copy_summary_row(low,header+16,['Low Total'])
    copy_summary_row(final,header+17,['Final Totals'])
    for col in ('C','E'):
        sheet[f'{col}{subtotal}']=f'=IF(COUNT({col}{first}:{col}{last})=0,"",SUM({col}{first}:{col}{last}))'
        sheet[f'{col}{final}']=f'=IF({col}{subtotal}="","",{col}{subtotal})'
    for col in ('B','D'):
        sheet[f'{col}{low}']=f'=IF(COUNT({col}{first}:{col}{last})=0,"",SUM({col}{first}:{col}{last}))'


def fit_columns(sheet,overage):
    """Fit each view to its own text, allowing old/template widths to shrink."""
    for target in (sheet,overage):
        for col in range(1,16):
            letter=target.cell(1,col).column_letter
            target.column_dimensions[letter].width=fitted_width(
                [row[0] for row in target.iter_rows(min_col=col,max_col=col)])
            target.column_dimensions[letter].bestFit=True


def generate(template, data, output, final=False, test_generation=False, *, _legacy_layout=False):
    if final and test_generation:
        raise ValueError('Test generation cannot be combined with final race assignments')
    races, horses, entries = prepare(data, final)
    output = Path(output)
    if output.resolve() == Path(template).resolve():
        raise ValueError('Output must differ from template')
    if output.exists():
        raise FileExistsError(f'Output already exists: {output}')
    original = load_workbook(template)
    source = original['2024']
    if [source.cell(1, c).value for c in range(1, 15)] != LEGACY_HEADERS[:14]:
        raise ValueError('Template 2024 headers differ from expected layout')
    result = Workbook()
    if data.get('demo'):
        result.properties.title=f'Breeders Cup {data["year"]} - DEMO'
        result.properties.description=data['demo']['notice']
    sheet = result.active
    sheet.title = str(data['year'])
    overage = result.create_sheet('Overage')
    for target in (sheet, overage):
        for col in 'ABCDEFGHIJKLMN':
            target.column_dimensions[col].width = source.column_dimensions[col].width
        target.column_dimensions['O'].width = 14
        target.sheet_view.zoomScale = 100
        target.freeze_panes = 'F2'
        target.sheet_properties.pageSetUpPr = copy(source.sheet_properties.pageSetUpPr)
        target.page_setup = copy(source.page_setup)
        target.print_title_rows = '1:1'
        styled_row(target, 1, LEGACY_HEADERS, source, 1)
        if data.get('demo'):
            text(target['A1'],'BC Race (DEMO)')
    row = 2
    displayed_entries=0
    for race in races:
        matching = sorted((e for e in entries if e['race_id'] == race['id']),
                          key=lambda e: post_order(e, horses[e['horse_id']]))
        if test_generation:
            matching=matching[:13]
        displayed_entries+=len(matching)
        when = date.fromisoformat(race['date'])
        label = f"Race {race['number']}" if race['number'] is not None else 'Race # pending'
        styled_row(sheet, row, [when.strftime('%A %m-%d-%Y'), label] + [None]*13, source, 4)
        row += 1
        for entry in matching:
            horse = horses[entry['horse_id']]
            merged = {**horse, **{k:v for k,v in entry.items() if k in FIELDS and k != 'name'}}
            merged['pick'] = None  # DE and H are personal picks, entered manually in Excel.
            merged['other_ancestors'] = famous_ancestors(horse, FAMOUS_HORSES)
            values = [race['name']] + [merged.get(f) for f in FIELDS]
            if merged.get('date_of_birth'):
                values[-1] = date.fromisoformat(merged['date_of_birth'])
            styled_row(sheet, row, values, source, 5)
            text(sheet.cell(row,19),horse['id'])
            text(sheet.cell(row,20),race['id'])
            row += 1
    if races:
        results_summary(sheet,row,races,source)
    # Overage is this year's complete research registry, one row per horse,
    # independent of whether race nominations or a schedule have been published.
    import json
    race_names = {r['id']: r['name'] for r in races}
    for r, (hid, horse) in enumerate(sorted(horses.items(), key=lambda item:
                                          (key_name(item[1]['name']), item[0])), 2):
        active = [e for e in entries if e['horse_id'] == hid]
        merged = dict(horse)
        if len(active) == 1:
            merged.update({k:v for k,v in active[0].items() if k in FIELDS and k != 'name'})
        else:
            # Posts, odds, results and engagements belong to a specific race.
            for field in ('post', 'result', 'odds'):
                merged[field] = None
            if active and all(e.get('scratched') or str(e.get('post','')).upper()=='SCR' for e in active):
                merged['odds']='SCR'
            riders={e.get('jockey') or horse.get('jockey') for e in active}
            riders.discard(None)
            merged['jockey']=next(iter(riders)) if len(riders)==1 else None
        merged['pick'] = None
        merged['other_ancestors'] = famous_ancestors(horse, FAMOUS_HORSES)
        values = ['; '.join(race_names[e['race_id']] for e in active) or None] + [merged.get(f) for f in FIELDS]
        if merged.get('date_of_birth'):
            values[-1] = date.fromisoformat(merged['date_of_birth'])
        styled_row(overage, r, values, source, 5)
        text(overage.cell(r,19),hid)
        for col, value in enumerate((horse['source_url'], horse.get('retrieved_at'),
                                     json.dumps(horse.get('field_sources', {}))), 16):
            text(overage.cell(r, col), value)
    for col, label in enumerate(('Source URL', 'Retrieved at', 'Field sources (JSON)'), 16):
        cell = overage.cell(1, col)
        text(cell, label)
        cell.font = copy(overage['A1'].font)
        cell.fill = copy(overage['A1'].fill)
    for col in ('P', 'Q', 'R'):
        overage.column_dimensions[col].width = 40
        overage.column_dimensions[col].hidden = True
    overage.auto_filter.ref = f'A1:T{max(1, len(horses)+1)}'
    for target in (sheet, overage):
        text(target['S1'],'Horse ID')
        text(target['T1'],'Race ID')
        target.column_dimensions['S'].hidden=True
        target.column_dimensions['T'].hidden=True
        target.print_area = f'A1:O{target.max_row}'
    # Excel's bestFit flag does not calculate a stored width. Estimate the text
    # width at the retained nine-point template font and save explicit widths.
    if not _legacy_layout:
        arrange_columns(result)
    fit_columns(sheet,overage)
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation also protects against another run creating the same output.
    with output.open('xb') as stream:
        result.save(stream)
    original.close()
    return displayed_entries


def key_name(name):
    return name.strip().casefold()
